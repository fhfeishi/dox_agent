"""K8: offline Word (.docx) parsing into searchable text."""

from docx import Document as DocxDocument

from src.agent.config import Settings
from src.agent.corpora import SOURCE_SUFFIXES
from src.parsers import parse_file


def test_docx_paragraphs_and_tables_become_text(tmp_path):
    path = tmp_path / "报告.docx"
    document = DocxDocument()
    document.add_paragraph("第一段正文")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "编号"
    document.save(path)

    parsed = parse_file(path, Settings(_env_file=None))
    assert parsed.kind == "text"
    assert parsed.parser.startswith("python-docx/")
    assert "第一段正文" in parsed.pages[0].text
    assert "项目 | 编号" in parsed.pages[0].text
    assert ".docx" in SOURCE_SUFFIXES
