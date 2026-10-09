from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .presentation import present_report

STATUS={'pass':'通过','issue':'发现问题','warning':'建议调整','pending':'待核实','na':'不适用','advisory':'评议建议'}

def export_report(run,path:Path):
    run=present_report(run)
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=sec.bottom_margin=Cm(2); sec.left_margin=sec.right_margin=Cm(2.3)
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        style=doc.styles[name]; style.font.name='Microsoft YaHei'; style._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        style.font.color.rgb=RGBColor.from_string('172B40')
    doc.styles['Normal'].font.size=Pt(10.5); doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Normal'].paragraph_format.line_spacing=1.25
    doc.styles['Title'].font.size=Pt(23); doc.styles['Title'].font.color.rgb=RGBColor(0,0,0)
    doc.styles['Heading 1'].font.size=Pt(16); doc.styles['Heading 2'].font.size=Pt(12)
    doc.add_paragraph("专业审查报告" if run["request"].get("kind") == "professional" else "形式审查报告", "Title")
    doc.add_paragraph(run['document']['metadata']['title'])
    counts=run['summary']; doc.add_paragraph(f'本次预审发现 {counts["issue"]} 项明确问题、{counts["warning"]} 项建议调整、{counts["pending"]} 项待核实。以下为内部预审结果，尚未覆盖的资格与政策要求仍需人工审核。')
    rows=[('审查模板',f'{run["rule"]["name"]} · v{run["rule"]["version"]}'),('生成时间',run['created_at'])]
    if run['request'].get('kind')=='professional':
        rows.append(('评价方式',f'历史评价 · 评价时点 {run["request"]["cutoff_date"]}' if run['request']['mode']=='historical' else '更新建议'))
    for label,value in rows: doc.add_paragraph(f'{label}：{value}')
    doc.add_heading('审核范围',1); doc.add_paragraph(run['rule']['scope_note'] or '依据所选模板开展审查，结果表示是否符合本模板要求。')
    revised=[c for c in run['rule']['checks'] if c.get('enabled',True) and c.get('revised')]
    if revised:
        doc.add_heading('模板人工修订',2)
        for c in revised: doc.add_paragraph(f'{c["title"]}：原文要求“{c["original"].get("requirement","")}”；本模板人工修订为“{c["requirement"]}”，按修订后要求执行。')
    doc.add_paragraph('PDF 使用文件物理页码；DOCX 使用文本块编号。报告固定保存当次规则、文档信息与证据快照。')
    reader=run.get('reader_report',{})
    doc.add_heading('审核结论与修改清单',1)
    doc.add_paragraph(reader.get('summary','请查看下方逐项结果。'))
    for heading,key in [('一、规范方面需要处理的问题','issues'),('二、研究方案的具体修改建议','technical'),('三、待补充材料与待核实事项','pending')]:
        doc.add_heading(heading,2)
        items=reader.get(key,[])
        if not items: doc.add_paragraph('本次未列出此类事项，不代表未核验的内容已经通过。')
        for index,item in enumerate(items,1):
            doc.add_paragraph(f'{index}. {item["title"]}（{STATUS.get(item["status"],item["status"])}）')
            doc.add_paragraph('建议：'+item.get('suggestion','请结合原文核查。'))
    if reader.get('missing_documents'):
        doc.add_heading('尚未提供的依据文件',2)
        for index,name in enumerate(reader['missing_documents'],1): doc.add_paragraph(f'{index}. {name}')
    if run['request'].get('kind') != 'professional': doc.add_heading('规范性审核',1)
    for f in run['findings']:
        doc.add_heading(f'{STATUS.get(f["status"],f["status"])}  {f["title"]}',2)
        doc.add_paragraph(f['detail'])
        if f.get('page'):
            location = run.get('source_locations', {}).get(str(f['page']))
            doc.add_paragraph(f"原文位置：{location['filename']} · 第 {location['source_page']} " + ('页' if location['kind'] == 'pdf' else '文本块') if location else f"原文位置：{f['page']}")
        doc.add_paragraph('判断依据：'+f['source'])
        if f.get('coverage')=='partial': doc.add_paragraph('读取覆盖：部分内容（如图片文字）未读取，完整结果未能核对。')
        if f.get('mapping') and f['mapping']!='confirmed': doc.add_paragraph('章节范围：自动识别或未定位，尚未人工确认。')
        if f.get('suggestion'): doc.add_paragraph('修改建议：'+f['suggestion'])
        for calc in f.get('calculations',[]): doc.add_paragraph(f'{calc["expression"]}；复算 {calc["computed_yuan"]:g} 元；'+('一致' if calc['ok'] else '不一致'))
    if run['request'].get('kind') != 'formal':
        doc.add_heading('技术分析与修改建议',1)
        for row in run['technical'].get('coverage',[]):
            if row['state']=='missing': doc.add_paragraph(f'未完成评价：{row["title"]}（本次模板启用，模型未给出有效意见）')
    tech=run['technical']; doc.add_paragraph(tech['summary'])
    if tech.get('error'): doc.add_paragraph('模型评议未完成：'+tech['error'])
    for card in tech['cards']:
        doc.add_heading(card['title'],2); doc.add_paragraph(card['assessment'])
        if card.get('quote'): doc.add_paragraph('原文摘录：'+card['quote'])
        if card.get('page'):
            location = run.get('source_locations', {}).get(str(card['page']))
            doc.add_paragraph(f"原文位置：{location['filename']} · {location['source_page']}" if location else f"原文位置：{card['page']}")
        doc.add_paragraph('建议：'+card['suggestion']); doc.add_paragraph('评议方式：'+card['basis'])
        if card['evidence_ids']: doc.add_paragraph('关联证据：'+', '.join(card['evidence_ids']))
    doc.add_heading('参考证据',1)
    for e in tech['evidence']:
        doc.add_heading(e['title'],2); doc.add_paragraph(f'{e["id"]} · {e.get("published") or "发表时间未明确"} · {e.get("origin") or e.get("kind", "参考资料")}')
        doc.add_paragraph(e['summary']); doc.add_paragraph(e['url'])
    if not tech['evidence']: doc.add_paragraph('本次未匹配到截止日期前的相关证据，不能据此确认原创性。')
    if run.get('audit'):
        audit=run['audit']; doc.add_heading('附录：审核范围与执行情况',1)
        unit='页' if run['document']['kind']=='pdf' else '文本块'
        doc.add_paragraph(f'服务：{audit["provider"]}；模型：{audit["model"]}；调用次数：{len(audit["calls"])}；已读取非空{unit}数：{len(audit["coverage_pages"])}')
        if audit.get('verification'):
            v=audit['verification']; doc.add_paragraph(f'最终复核：{v["completed"]}/{v["total"]} 项完成；{v["pending"]} 项未完成最终复核。复核完成不等于结论通过。')
        doc.add_paragraph('逐项待核实状态与证据见正文；原始模型输出和调用记录保存在本次审核的 JSON 文件中。')
    foot=sec.footer.paragraphs[0]; foot.alignment=2
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); foot._p.append(field)
    path.parent.mkdir(parents=True,exist_ok=True); doc.save(path)
    return path
