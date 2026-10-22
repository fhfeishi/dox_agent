"""One real HTTP/material pipeline distinguishes formal checks from professional review.

Only model responses are controlled; uploads, DOCX parsing, rule publication, persistence,
execution dispatch, calculations and result recovery use real code and temporary storage.
"""

import time
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from src.agent.config import Settings
from src.main import create_app
from src.review import routes, storage


class FixtureReviewClient:
    phases = []

    def __init__(self):
        self.settings = {"model": "fixture", "max_chars": 240000, "chunk": 12000, "max_calls": 32}
        self.calls = []

    def ask(self, phase, payload):
        self.phases.append(phase)
        if phase == "01_extract":
            return {"overview": "申请正文", "items": [{"source_id": payload["sources"][0]["source_id"]}]}
        if phase == "03_compliance":
            return {
                "findings": [
                    {
                        "title": c["title"],
                        "detail": "本项需要原文及计算依据核对。",
                        "suggestion": "检查原文",
                        "group": "guideline",
                        "status": "pending",
                        "clause_ids": [c["id"]],
                        "tool_ids": [],
                    }
                    for c in payload["clauses"]
                ]
            }
        if phase == "04_technical":
            return {
                "summary": "现有资料不足以判断原创。",
                "cards": [
                    # The last question answered first: the report must still follow template order.
                    *([{"check_id": payload["questions"][-1]["check_id"], "title": "差异待说明",
                        "assessment": "需与相近项目逐项对照", "suggestion": "列出对照项目",
                        "source_id": payload["facts"][0]["source_id"], "evidence_ids": []}]
                      if len(payload["questions"]) > 1 else []),
                    {
                        "check_id": payload["questions"][0]["check_id"],
                        "title": "创新依据待补充",
                        "assessment": "未检索到相近项目不等于原创",
                        "suggestion": "补充比较基线",
                        # Real models join several passages into one field; the first valid one anchors the card.
                        "source_id": "unknown-id; " + payload["facts"][0]["source_id"],
                        "evidence_ids": [],
                    },
                ],
            }
        if phase == "05_verify":
            return {
                "summary": "已核对",
                "actions": [
                    {"target_id": ident, "decision": "downgrade", "reason": "模拟模型降级请求"}
                    for ident in payload["expected_ids"]
                ],
            }
        raise AssertionError(phase)


def test_two_material_tasks_are_independent_and_keep_program_findings(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "DATA", tmp_path / "review")
    FixtureReviewClient.phases.clear()
    calls = FixtureReviewClient.phases
    monkeypatch.setattr(routes, "Client", FixtureReviewClient)
    monkeypatch.setattr(routes, "model_configured", lambda: True)
    monkeypatch.setattr(
        routes, "start_metadata", lambda ident: routes.public_doc(storage.get("document", ident))
    )
    settings = Settings(
        _env_file=None,
        state_dir=tmp_path / "state",
        corpora_root=tmp_path / "corpora",
    )
    (settings.corpora_root / "fixture" / "source").mkdir(parents=True)
    app = create_app(settings)
    doc = Document()
    doc.styles["Normal"].element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    doc.add_heading("技术方案", 1)
    doc.add_paragraph("本方案申请直接费用为550万元，需要临床数据验证，不能据此宣称原创。" * 12)
    doc.add_paragraph().add_run("黑体四字").font.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "黑体")
    data = BytesIO()
    doc.save(data)
    with TestClient(app) as client:
        uploaded = client.post("/api/review/documents", files={"file": ("proposal.docx", data.getvalue())})
        assert uploaded.status_code == 200, uploaded.text
        proposal = uploaded.json()
        assert (
            client.patch(
                "/api/review/documents/" + proposal["id"],
                json={"title": "测试申请", "budget": 550, "organization_count": 6, "submitted_at": "2026-10-01T09:00:00+08:00", "year": 2026},
            ).status_code
            == 200
        )
        check = {
            "id": "budget-limit",
            "title": "申请直接费用上限",
            "category": "预算",
            "requirement": "申请直接费用不超过500万元",
            "applicability": "本批申请",
            "strength": "hard",
            "method": "calculation",
            "needed_materials": ["信息表"],
            "source": "示例指南原文",
            "source_id": "clause-1",
            "guideline_id": "g1",
            "quote": "申请直接费用不超过500万元",
            "page": 1,
            "execution": {"field": "budget", "operator": "le", "value": "500", "unit": "万元"},
        }
        checks = [check]
        for ident, field, value, unit, operator in [
            ("organization-limit", "organization_count", "5", "家", "le"),
            ("submission-deadline", "submitted_at", "2026-10-10T17:00:00+08:00", "ISO8601", "before"),
        ]:
            checks.append({**check, "id": ident, "title": ident,
                           "execution": {"field": field, "operator": operator, "value": value, "unit": unit}})
        checks[1].update(title="申报单位数量上限", requirement="申报单位不超过5家", quote="申报单位不超过5家")
        checks.append({**check, "id": "body-font", "title": "正文字体", "needed_materials": [],
                       "execution": {"field": "docx_font", "operator": "eq", "value": "宋体", "unit": "字体"}})
        checks[2].update(title="申报截止时间", requirement="申报时间应早于2026年10月10日17时北京时间",
                         quote="申报时间应早于2026年10月10日17时北京时间")
        assert (
            client.post(
                "/api/review/rules",
                json={
                    "id": "fixture",
                    "kind": "formal",
                    "name": "测试要求",
                    "fund": "示例",
                    "year": 2026,
                    "scope_note": "测试规则，不是正式基金标准",
                    "checks": checks,
                    "confirmed": True,
                },
            ).status_code
            == 200
        )

        def wait(run_id):
            for _ in range(100):
                run = client.get("/api/review/runs/" + run_id).json()
                if run["status"] not in ("queued", "running"):
                    return run
                time.sleep(0.02)
            raise AssertionError("run did not finish")

        formal = client.post(
            "/api/review/runs",
            json={
                "document_id": proposal["id"],
                "rule_id": "fixture",
                "rule_version": 1,
                "kind": "formal",
            },
        ).json()
        formal = wait(formal["id"])
        assert formal["status"] == "completed", formal
        assert "04_technical" not in calls
        assert not any(
            f["status"] == "pass" for f in formal["findings"] if "budget-limit" in f.get("clause_ids", [])
        )
        assert any("信息表" in f["title"] for f in formal["findings"])
        sheet = client.post(
            "/api/review/documents", files={"file": ("information.docx", data.getvalue())}
        ).json()
        client.patch(
            "/api/review/documents/" + sheet["id"], json={"title": "测试申请", "budget": 550, "organization_count": 6, "submitted_at": "2026-10-01T09:00:00+08:00", "year": 2026}
        )
        with_sheet = client.post(
            "/api/review/runs",
            json={
                "document_id": proposal["id"],
                "information_sheet_id": sheet["id"],
                "rule_id": "fixture",
                "rule_version": 1,
                "kind": "formal",
            },
        ).json()
        with_sheet = wait(with_sheet["id"])
        assert with_sheet["status"] == "completed", with_sheet
        assert any(f["status"] == "issue" and f.get("origin") == "program" for f in with_sheet["findings"])
        computed = {item["id"]: item["status"] for item in with_sheet["calculation_findings"]}
        assert computed["check-organization-limit"] == "issue"
        assert computed["check-submission-deadline"] == "pass"
        font = next(f for f in with_sheet["calculation_findings"] if f["id"] == "check-body-font")
        assert font["status"] == "issue" and "4 字不是宋体，0 字" in font["detail"]
        assert all(
            item["file_id"] in (proposal["id"], sheet["id"])
            for item in with_sheet["source_locations"].values()
        )

        before = len(calls)
        cid = client.get("/api/corpora").json()[0]["id"]
        professional = client.post(
            "/api/review/runs",
            json={
                "document_id": proposal["id"],
                "kind": "professional",
                "rule_id": "builtin-professional",
                "rule_version": 1,
                "corpus_ids": [cid],
                "cutoff_date": "2026-09-30",
            },
        ).json()
        professional = wait(professional["id"])
        assert professional["status"] == "completed", professional
        assert "03_compliance" not in calls[before:]
        assert professional["request"]["kind"] == "professional"
        assert professional["technical"]["cards"] and not professional["findings"]
        assert professional["technical"]["cards"][0]["quote"]
        assert client.get("/api/review/runs/" + professional["id"] + "/export/docx").status_code == 200
        assert client.get("/api/review/runs/" + with_sheet["id"] + "/export/docx").status_code == 200
        assert professional["technical"]["local_matches"] == 0
        assert "资料" in professional["reader_report"]["summary"] or professional["technical"]["summary"]


def test_generic_template_sections_versions_and_internal_review(tmp_path, monkeypatch):
    """Builtin six-item template: section ranges, boundaries, origins, and evidence-free review."""
    monkeypatch.setattr(storage, "DATA", tmp_path / "review")
    monkeypatch.setattr(routes, "Client", FixtureReviewClient)
    monkeypatch.setattr(routes, "model_configured", lambda: True)
    monkeypatch.setattr(routes, "start_metadata", lambda ident: routes.public_doc(storage.get("document", ident)))
    settings = Settings(_env_file=None, state_dir=tmp_path / "state",
                        corpora_root=tmp_path / "corpora")
    (settings.corpora_root / "fixture" / "source").mkdir(parents=True)
    with TestClient(create_app(settings)) as client:
        def run(body):
            created = client.post("/api/review/runs", json=body)
            assert created.status_code == 200, created.text
            for _ in range(200):
                result = client.get("/api/review/runs/" + created.json()["id"]).json()
                if result["status"] not in ("queued", "running"):
                    assert result["status"] == "completed", result
                    return result
                time.sleep(.02)
            raise AssertionError("run did not finish")

        template = client.get("/api/review/rules").json()
        builtin = next(r for r in template if r.get("id") == "builtin-formal")
        assert client.post("/api/review/rules", json=builtin).status_code == 403
        body = ("# 项目摘要\n" + "甲" * 300 + "\n# 项目背景\n" + "乙" * 201 + "\n# 项目方案\n" + "丙" * 1990
                + "\n## 技术路线\n" + "丁" * 10 + "\n# 项目进度计划\n2026年10月至2027年9月完成原型；2028.10.1 验收。")
        doc = client.post("/api/review/documents", files={"file": ("proposal.md", body.encode())}).json()
        assert client.patch("/api/review/documents/" + doc["id"], json={
            "title": "通用模版申请", "organizations": ["甲大学", "甲大学 ", "乙公司", "丙研究院"],
            "total_budget": 300.01, "plan_start": "2026-10-01", "plan_end": "2028-09-30"}).status_code == 200
        formal = {"document_id": doc["id"], "kind": "formal", "rule_id": builtin["id"], "rule_version": builtin["version"]}

        # A unique, completely read range is located automatically and judged without a confirmation click.
        auto = {f["id"]: f for f in run(formal)["calculation_findings"]}
        assert auto["check-abstract"]["status"] == "pass" and auto["check-abstract"]["actual"] == 300
        assert auto["check-abstract"]["mapping"] == "automatic"

        prep = client.get(f"/api/review/documents/{doc['id']}/mappings", params={"rule_id": builtin["id"]}).json()
        assert [m["status"] for m in prep["mappings"]] == ["found"] * 4
        update = {"rule_id": builtin["id"], "rule_version": builtin["version"], "revision": 0, "mappings": prep["mappings"]}
        assert client.put(f"/api/review/documents/{doc['id']}/mappings", json=update).status_code == 200
        assert client.put(f"/api/review/documents/{doc['id']}/mappings", json=update).status_code == 409
        results = {f["id"]: f for f in run(formal)["calculation_findings"]}
        assert results["check-abstract"]["status"] == "pass" and results["check-abstract"]["actual"] == 300
        assert results["check-background"]["status"] == "issue"
        assert results["check-plan"]["status"] == "pass" and results["check-plan"]["actual"] == 2000  # Subsection counted, titles not.
        assert results["check-organizations"]["status"] == "pass" and results["check-organizations"]["actual"] == 3
        assert results["check-budget"]["status"] == "issue"
        assert results["check-schedule"]["status"] == "issue" and "2028.10.1" in results["check-schedule"]["detail"]

        custom = {**builtin, "id": "my-template", "name": "我的模版", "source_kind": "manual", "version": 1}
        custom["checks"][1]["execution"]["value"] = "500"
        custom["checks"][1]["requirement"] = "项目摘要不超过500字符"
        created = client.post("/api/review/rules", json=custom)
        assert created.status_code == 200, created.text
        assert created.json()["checks"][1]["revised"]
        assert created.json()["checks"][1]["original"]["execution"]["value"] == "300"
        assert client.post("/api/review/rules", json={**custom, "kind": "professional"}).status_code == 422
        disabled = {**custom, "checks": [{**c, "enabled": False} for c in custom["checks"]]}
        assert client.post("/api/review/rules", json=disabled).status_code == 422

        # Professional review compares against exactly one library; stock literature is unread unless chosen.
        client.post("/api/review/evidence", json={"title": "库存文献", "published": "2020-01-01",
                                                  "url": "https://example.org/a", "summary": "与本申请无关的库存文献摘要内容。"})
        body = {"document_id": doc["id"], "kind": "professional", "rule_id": "builtin-professional", "rule_version": 1}
        assert client.post("/api/review/runs", json=body).status_code == 422
        assert client.post("/api/review/runs", json={**body, "corpus_ids": ["no-such-library"]}).status_code == 422
        library = client.get("/api/corpora").json()[0]["id"]
        professional = run({**body, "corpus_ids": [library]})
        assert professional["request"]["reference_count"] == 3
        assert professional["technical"]["evidence"] == []
        coverage = {row["check_id"]: row["state"] for row in professional["technical"]["coverage"]}
        assert coverage["objective"] == "evaluated" and coverage["risk"] == "missing"
        # Report order (query 2026-1010 0914) and opinions follow it whatever order the model used.
        order = [row["check_id"] for row in professional["technical"]["coverage"]]
        assert order == ["objective", "necessity", "feasibility", "advancement", "risk", "comparison"]
        ranks = [order.index(card["check_id"]) for card in professional["technical"]["cards"]]
        assert len(ranks) == 2 and ranks == sorted(ranks), ranks
        assert any("仅检索到 0 篇相近报告" in note for note in professional["audit"]["limitations"])


def test_metadata_extraction_and_confirmation_feed_checks_safely(tmp_path, monkeypatch):
    """Real extraction pipeline; the model only returns candidates that cite supplied sources."""
    monkeypatch.setattr(storage, "DATA", tmp_path / "review")

    class MetadataClient(FixtureReviewClient):
        def ask(self, phase, payload):
            if phase != "00_metadata":
                return super().ask(phase, payload)
            wanted = [("title", "智能康复机器人研究"), ("organizations", "甲大学"), ("organizations", "乙医院"),
                      ("total_budget", 280), ("plan_start", "2027-01"), ("plan_end", "2028-12")]
            rows = []
            for field, value in wanted:
                quote = {"total_budget": "280万元", "plan_start": "2027年1月", "plan_end": "2028年12月"}.get(field, str(value))
                source = next((s for s in payload["sources"] if quote in s["text"]), None)
                if source:
                    rows.append({"field": field, "value": value, "source_id": source["source_id"], "reason": "原文"})
            return {"candidates": rows, "notes": []}

    monkeypatch.setattr(routes, "Client", MetadataClient)
    monkeypatch.setattr(routes, "model_configured", lambda: True)
    monkeypatch.setattr("src.review.metadata.Client", MetadataClient)
    settings = Settings(_env_file=None, state_dir=tmp_path / "state", corpora_root=tmp_path / "corpora")
    body = ("# 项目基本信息\n项目名称：智能康复机器人研究\n申报单位：甲大学；乙医院\n项目总预算：280万元\n"
            "执行期限：2027年1月至2028年12月\n# 项目摘要\n" + "康复训练研究内容。" * 20)
    with TestClient(create_app(settings)) as client:
        def settle(ident):
            for _ in range(200):
                doc = client.get("/api/review/documents/" + ident).json()
                if doc["metadata_extraction"].get("status") != "running":
                    return doc
                time.sleep(.02)
            raise AssertionError("extraction did not finish")

        doc = settle(client.post("/api/review/documents", files={"file": ("p.md", body.encode())}).json()["id"])
        meta = doc["metadata"]
        assert (meta["title"], meta["organizations"], meta["total_budget"], meta["plan_start"], meta["plan_end"]) == (
            "智能康复机器人研究", ["甲大学", "乙医院"], 280, "2027-01", "2028-12")
        # A confirmed correction survives re-extraction; the model value stays a suggestion.
        assert client.patch("/api/review/documents/" + doc["id"], json={**meta, "total_budget": 250}).status_code == 200
        client.post(f"/api/review/documents/{doc['id']}/extract")
        doc = settle(doc["id"])
        assert doc["metadata"]["total_budget"] == 250 and doc["confirmed_at"]
        assert any("total_budget" in note for note in doc["metadata_extraction"]["notes"])

        # A confirmed sheet with no unit list does not erase the body's list; a differing budget conflicts.
        sheet = settle(client.post("/api/review/documents", files={"file": ("s.md", body.encode())}).json()["id"])
        client.patch("/api/review/documents/" + sheet["id"], json={"title": "信息表", "organizations": [], "total_budget": 300})
        rule = next(r for r in client.get("/api/review/rules").json() if r["id"] == "builtin-formal")

        def calc(document_id):
            created = client.post("/api/review/runs", json={"document_id": document_id, "information_sheet_id": sheet["id"],
                                                           "kind": "formal", "rule_id": rule["id"], "rule_version": rule["version"]})
            assert created.status_code == 200, created.text
            for _ in range(200):
                run = client.get("/api/review/runs/" + created.json()["id"]).json()
                if run["status"] not in ("queued", "running"):
                    return {f["id"]: f for f in run["calculation_findings"]}
                time.sleep(.02)
            raise AssertionError("run did not finish")

        results = calc(doc["id"])
        assert results["check-organizations"]["actual"] == 2
        assert results["check-budget"]["status"] == "pending" and "conflict-total_budget" in results
        # An unconfirmed body contributes nothing even when the sheet is confirmed.
        unconfirmed = settle(client.post("/api/review/documents", files={"file": ("u.md", body.encode())}).json()["id"])
        assert unconfirmed["metadata"]["organizations"] and not unconfirmed.get("confirmed_at")
        results = calc(unconfirmed["id"])
        assert results["check-organizations"]["status"] == "pending" and results["check-budget"]["actual"] == 300


def test_queued_reports_accept_three_tasks_and_auto_map_uploaded_sections(tmp_path, monkeypatch):
    import threading
    gate = threading.Event()

    class WaitingModel(FixtureReviewClient):
        def ask(self, phase, payload):
            if phase == "01_extract":
                assert gate.wait(8), "model fixture was not released"
            return super().ask(phase, payload)

    monkeypatch.setattr(storage, "DATA", tmp_path / "review")
    monkeypatch.setattr(routes, "Client", WaitingModel)
    monkeypatch.setattr(routes, "model_configured", lambda: True)
    monkeypatch.setattr(routes, "start_metadata", lambda ident: routes.public_doc(storage.get("document", ident)))
    settings = Settings(_env_file=None, state_dir=tmp_path / "state", corpora_root=tmp_path / "corpora")
    with TestClient(create_app(settings)) as client:
        body = "# 项目摘要\n" + "甲" * 300 + "\n# 项目背景\n" + "乙" * 201 + "\n# 项目方案\n" + "丙" * 100
        doc = client.post("/api/review/documents", files={"file": ("queue.md", body.encode())}).json()
        ids = []
        try:
            for _ in range(3):
                response = client.post("/api/review/runs", json={"document_id": doc["id"], "kind": "formal",
                                       "rule_id": "builtin-formal", "rule_version": 1})
                assert response.status_code == 200, response.text
                ids.append(response.json()["id"])
            last = client.get("/api/review/runs/" + ids[-1]).json()
            assert last["status"] == "queued"
            assert len(client.get("/api/review/runs").json()) == 3
        finally:
            gate.set()
        for ident in ids:
            for _ in range(250):
                result = client.get("/api/review/runs/" + ident).json()
                if result["status"] not in ("queued", "running"):
                    break
                time.sleep(.02)
            assert result["status"] == "completed", result
            measured = {f["id"]: f for f in result["calculation_findings"]}
            assert measured["check-abstract"]["status"] == "pass"
            assert measured["check-background"]["status"] == "issue"
            assert measured["check-abstract"]["mapping"] == "automatic"
        # Startup recovery distinguishes interrupted queue entries from live work.
        storage.save("run", {"id": "interrupted", "status": "queued"})
        storage.init()
        assert storage.get("run", "interrupted")["status"] == "failed"


def test_local_evidence_file_is_previewable_and_reaches_the_reviewer_as_text(tmp_path, monkeypatch):
    """query 2026-1013 1520: restricted networks rely on uploaded local literature."""
    seen = []

    class RecordingClient(FixtureReviewClient):
        def ask(self, phase, payload):
            if phase == "04_technical":
                seen.append(payload["evidence"])
            return super().ask(phase, payload)

    monkeypatch.setattr(storage, "DATA", tmp_path / "review")
    monkeypatch.setattr(routes, "Client", RecordingClient)
    monkeypatch.setattr(routes, "model_configured", lambda: True)
    monkeypatch.setattr(routes, "start_metadata", lambda ident: routes.public_doc(storage.get("document", ident)))
    settings = Settings(_env_file=None, state_dir=tmp_path / "state", corpora_root=tmp_path / "corpora")
    (settings.corpora_root / "fixture" / "source").mkdir(parents=True)
    paper = "# 柔性装配综述\n" + "多模态感知驱动的柔性装配方法在工业场景中已有验证，精度达到亚毫米级。" * 6
    with TestClient(create_app(settings)) as client:
        files = {"file": ("综述.md", paper.encode())}
        assert client.post("/api/review/evidence/upload", files=files, data={"published": "2024-13-01"}).status_code == 422
        assert client.post("/api/review/evidence/upload", files={"file": ("short.md", "# 空".encode())}).status_code == 400
        item = client.post("/api/review/evidence/upload", files=files, data={"published": "2024-05-01"}).json()
        assert (item["title"], item["published"], item["filename"]) == ("综述", "2024-05-01", "综述.md")
        assert "pages" not in item and "path" not in item and "亚毫米级" in item["summary"]
        listed = client.get("/api/review/evidence").json()
        assert [e["id"] for e in listed] == [item["id"]] and "pages" not in listed[0]
        assert any("亚毫米级" in page["text"] for page in client.get(f"/api/review/evidence/{item['id']}/pages").json())
        assert client.get(f"/api/review/evidence/{item['id']}/file").content == paper.encode()

        body = "# 项目摘要\n" + "本项目研究柔性装配。" * 20 + "\n# 研究目标\n" + "实现多模态感知驱动的柔性装配。" * 20
        doc = client.post("/api/review/documents", files={"file": ("proposal.md", body.encode())}).json()
        library = client.get("/api/corpora").json()[0]["id"]
        created = client.post("/api/review/runs", json={"document_id": doc["id"], "kind": "professional", "rule_id": "builtin-professional",
                                                        "rule_version": 1, "corpus_ids": [library], "evidence_ids": [item["id"]]})
        assert created.status_code == 200, created.text
        for _ in range(250):
            if client.get("/api/review/runs/" + created.json()["id"]).json()["status"] not in ("queued", "running"):
                break
            time.sleep(.02)
    assert seen and [e.get("title") for e in seen[0]] == ["综述"]
    assert all("pages" not in e and "path" not in e for e in seen[0])


def test_review_templates_can_be_deleted_but_built_in_ones_stay(tmp_path, monkeypatch):
    """query 2026-1014 1737: 删除 takes a template out of the list; past reviews keep their version."""
    monkeypatch.setattr(storage, "DATA", tmp_path / "review")
    settings = Settings(_env_file=None, state_dir=tmp_path / "state", corpora_root=tmp_path / "corpora")
    with TestClient(create_app(settings)) as client:
        rules = client.get("/api/review/rules").json()
        builtin = next(rule for rule in rules if rule["id"].startswith("builtin-"))
        storage.save("rule", {**builtin, "id": "mine", "name": "我的模版", "source_kind": "manual"})
        assert client.delete("/api/review/rules/mine").json() == {"id": "mine", "archived": True}
        assert next(rule for rule in client.get("/api/review/rules").json() if rule["id"] == "mine")["archived"]
        assert client.delete(f"/api/review/rules/{builtin['id']}").status_code == 403
        assert client.delete("/api/review/rules/none").status_code == 404


def test_an_applied_template_is_what_new_reviews_start_with(tmp_path, monkeypatch):
    """query 2026-1021 1629: 应用 makes a template the starting choice for that review kind, 不应用 goes
    back to the built-in one, and a template that stops being usable stops being applied."""
    monkeypatch.setattr(storage, "DATA", tmp_path / "review")
    settings = Settings(_env_file=None, state_dir=tmp_path / "state", corpora_root=tmp_path / "corpora")
    with TestClient(create_app(settings)) as client:
        assert client.get("/api/review/applied").json() == {"formal": None, "professional": None}
        formal = next(rule for rule in client.get("/api/review/rules").json() if rule["id"] == "builtin-formal")
        storage.save("rule", {**formal, "id": "mine", "name": "按指南生成", "source_kind": "upload"})
        storage.save("rule", {**formal, "id": "draft", "name": "草稿", "source_kind": "upload", "confirmed": False})

        # Given a confirmed template, applying it makes it the formal starting choice only
        assert client.put("/api/review/applied", json={"kind": "formal", "rule_id": "mine"}).json() == \
            {"formal": "mine", "professional": None}
        # A draft, a template of the other kind and an unknown one are refused
        assert client.put("/api/review/applied", json={"kind": "formal", "rule_id": "draft"}).status_code == 422
        assert client.put("/api/review/applied", json={"kind": "professional", "rule_id": "mine"}).status_code == 422
        assert client.put("/api/review/applied", json={"kind": "formal", "rule_id": "none"}).status_code == 404

        # Deleting the applied template quietly returns the kind to the built-in one
        client.delete("/api/review/rules/mine")
        assert client.get("/api/review/applied").json()["formal"] is None

        # 不应用 clears the choice
        storage.save("rule", {**formal, "id": "again", "name": "再生成", "source_kind": "upload"})
        client.put("/api/review/applied", json={"kind": "formal", "rule_id": "again"})
        assert client.put("/api/review/applied", json={"kind": "formal", "rule_id": None}).json()["formal"] is None
