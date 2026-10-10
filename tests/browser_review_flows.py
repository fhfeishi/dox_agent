"""Real HTTP/storage/browser journey for project identity, external input and two review tasks.

Only the model is replaced: no remote model request is made. Materials, DOCX parsing,
rule publication, task dispatch, calculations, versioned artifacts and readers are real.
"""

import asyncio
import re
import json
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from urllib.parse import quote

from docx import Document as WordDocument
from playwright.async_api import async_playwright, expect
from pydantic import SecretStr

from src.agent.config import Settings
from src.agent.corpora import scan_corpora
from src.knowledge import Document, Knowledge, Page
from src.main import create_app
from src.review import routes, storage
from src.targets import DIMENSIONS, PROMPT_VERSION, SCHEMA_VERSION, _save
from tests.browser_artifacts_live import start_server
from tests.test_review_flows import FixtureReviewClient


async def main():
    import src.intelligence as intelligence

    class AnalysisModel:
        async def ainvoke(self, messages):
            return SimpleNamespace(
                content="# 验收情报报告\n\n## 技术背景\n外部材料提出的技术条件需要进一步核对 [1]。\n\n## 来源局限\n只依据本次材料，不据此判定先进性。"
            )

    intelligence.model_for = lambda settings: AnalysisModel()

    class SummaryModel:
        async def ainvoke(self, messages):
            return SimpleNamespace(content=json.dumps({
                "overview": "两个项目都把临床诊疗作为应用场景。",
                "common": [{"text": "都面向临床诊疗中的病灶识别", "refs": ["E1", "E2"]}],
                "differences": []}, ensure_ascii=False))

    import src.topic_summary as topic_summary
    topic_summary.model_for = lambda settings: SummaryModel()
    routes.Client = FixtureReviewClient
    routes.model_configured = lambda: True
    routes.start_metadata = lambda ident: routes.public_doc(storage.get("document", ident))
    with TemporaryDirectory(prefix="dox-business-browser-") as directory:
        root = Path(directory)
        corpus = root / "corpora" / "自然科学基金-验收库"
        source = corpus / "source"
        source.mkdir(parents=True)
        settings = Settings(
            _env_file=None,
            corpora_root=root / "corpora",
            state_dir=root / "state",
            model_api_key=SecretStr("offline"),
            embedding_path="",
        )
        storage.DATA = root / "review"
        knowledge = Knowledge(corpus / "datadb" / "knowledge.sqlite3", settings=settings, source_root=source)
        documents = []
        for filename, title, money in [
            ("2021_2025_A01_甲_影像报告.md", "影像项目", 100),
            ("2021_2025_A01_甲_影像摘要.md", "影像项目", 120),
            ("2023_2025_A02_乙_病理报告.md", "病理项目", 200),
        ]:
            body = f"# 国家自然科学基金\n项目名称：{title}\n直接费用：{money}（万元）\n临床诊疗场景，采用深度学习解决病灶识别困难，形成原型系统。"
            path = source / filename
            path.write_text(body)
            documents.append(
                knowledge.put(
                    Document(
                        title=title,
                        origin=str(path),
                        kind="text",
                        parser="markdown",
                        pages=[Page(number=1, text=body)],
                        markdown=body,
                    )
                )
            )
        info = scan_corpora(settings)[0]
        for saved in (documents[0], documents[2]):
            _save(
                info,
                {
                    "doc_id": saved["doc_id"],
                    "version": saved["version"],
                    "schema_version": SCHEMA_VERSION,
                    "prompt_version": PROMPT_VERSION,
                    "process": {"status": "已完成", "coverage": {"processed": 1, "total": 1}},
                    "facets": [
                        {
                            "key": dim,
                            "state": "has",
                            "items": [
                                {
                                    "id": dim,
                                    "name": ["临床诊疗", "识别困难", "深度学习", "原型系统"][i],
                                    "desc": "保留原文条件与出处。",
                                    **({"status": "已取得"} if dim == "成果" else {}),
                                    "evidence": [
                                        {
                                            "version": saved["version"],
                                            "quote": "临床诊疗场景",
                                            "locator": {"basis": "parsed_text", "page": 1},
                                        }
                                    ],
                                }
                            ],
                        }
                        for i, dim in enumerate(DIMENSIONS)
                    ],
                    "relations": [],
                },
            )
        app = create_app(settings, knowledge)
        origin, server, thread, listener = start_server(app)
        screenshots = Path(".logsdev/verification/business-20261001")
        screenshots.mkdir(parents=True, exist_ok=True)
        try:
            async with async_playwright() as playwright:
                browser = await playwright.chromium.launch(headless=True)
                page = await browser.new_page(viewport={"width": 1440, "height": 1000})
                errors = []
                page.on("pageerror", lambda error: errors.append(str(error)))
                await page.goto(origin + "/library/" + quote(info.id) + "/targets/four-facets")
                overview = page.get_by_label("项目归纳", exact=True)
                await expect(
                    overview.get_by_text("本库已识别项目", exact=True).locator("..")
                ).to_contain_text("2")
                await expect(
                    page.get_by_role("navigation").first.get_by_role(
                        "button", name="Prompt / Skill", exact=True
                    )
                ).to_have_count(0)
                # 四维浏览 lands on the hierarchy; 主题归纳 is gone, topic upkeep lives in 管理资料.
                await expect(overview.get_by_label("场景层级", exact=True)).to_contain_text("尚未生成")
                await expect(overview.get_by_label("更多视图").get_by_role("button")).to_have_text(["年度分布", "项目关系"])
                await page.get_by_role("button", name="管理资料", exact=True).first.click()
                admin = page.get_by_label("四维条目整理", exact=True)
                await admin.get_by_text("四维条目整理（合并同义、改名、撤销，查看条目归纳）", exact=True).click()
                await admin.get_by_text("整理场景条目（合并同义、改名、撤销）", exact=True).click()
                await admin.get_by_label("选择条目：临床诊疗", exact=True).check()
                await admin.get_by_label("整理后的条目名称", exact=True).fill("临床诊疗与辅助诊断")
                await admin.get_by_role("button", name="改名", exact=True).click()
                await expect(admin.get_by_text("原条目：临床诊疗", exact=True)).to_be_visible()
                await admin.get_by_role("button", name="撤销", exact=True).click()
                await expect(admin.get_by_text("已撤销", exact=True)).to_be_visible()
                await expect(admin.get_by_text("原条目：临床诊疗", exact=True)).to_have_count(0)
                await admin.get_by_label("查看条目归纳", exact=True).select_option("临床诊疗")
                digest = admin.get_by_label("主题归纳", exact=True)
                await expect(digest.get_by_text("尚未生成", exact=True)).to_be_visible()
                await digest.get_by_role("button", name="生成主题归纳", exact=True).click()
                await expect(digest.get_by_text("两个项目都把临床诊疗作为应用场景。", exact=True)).to_be_visible()
                await expect(digest.get_by_text("与当前资料一致", exact=True)).to_be_visible()
                await digest.get_by_text("依据 2 条 · 2 个项目", exact=True).click()
                await digest.screenshot(path=str(screenshots / "topic-summary.png"), animations="disabled")
                await digest.get_by_role("button", name="解析正文位置 · 查看原文 临床诊疗场景").first.click()
                await expect(page.get_by_role("complementary", name="检查器")).to_be_visible()
                await page.get_by_role("button", name="关闭检查器", exact=True).click()
                await page.get_by_role("button", name="四维浏览", exact=True).first.click()
                await page.get_by_role("button", name="年度分布", exact=True).click()
                await expect(overview.get_by_label("年度项目数量")).to_be_visible()
                await expect(overview.get_by_role("button", name="2021年：1个项目", exact=True)).to_be_visible()
                await expect(overview.get_by_role("button", name="2022年：0个项目", exact=True)).to_be_visible()
                await overview.get_by_role("button", name="2022年：0个项目", exact=True).click()
                await expect(overview.get_by_role("heading", name="关联项目 · 0", exact=True)).to_be_visible()
                await overview.get_by_role("button", name="重置筛选", exact=True).click()
                await page.get_by_role("button", name="时间与经费", exact=True).click()
                await expect(overview.get_by_role("img")).to_be_visible()
                await overview.get_by_role("button", name="病理项目", exact=True).click()
                await expect(page.get_by_label("项目详情", exact=True)).to_contain_text("批准号 A02")
                await page.get_by_role("button", name="查看此文件的原文依据", exact=False).first.click()
                await expect(page.get_by_role("complementary", name="检查器")).to_contain_text("四维信息")
                await page.get_by_role("button", name="返回上一预览").click()
                await expect(page.get_by_label("项目详情", exact=True)).to_contain_text("批准号 A02")
                await page.get_by_role("button", name="返回上一预览").click()
                await page.screenshot(
                    path=str(screenshots / "projects.png"), full_page=True, animations="disabled"
                )
                await page.get_by_role("button", name="项目关系", exact=True).click()
                graph = overview.get_by_label("项目关系图", exact=True)
                await expect(graph.locator("g[role=button]")).to_have_count(2)
                # Without a hierarchy only same-named technical topics can relate projects; 成果 is no relation type.
                await expect(graph.get_by_role("group", name="关联依据").get_by_role("button")).to_have_text(["场景", "问题", "技术"])
                await graph.get_by_role("button", name="技术", exact=True).click()
                await expect(graph.locator("g[role=button][data-degree='1']")).to_have_count(2)
                await graph.get_by_role("button", name="影像项目，2021，代码未知").click()
                detail = page.get_by_label("项目关系详情", exact=True)
                await expect(detail).to_contain_text("技术关联项目 · 1")
                await expect(detail).to_contain_text("共同技术：深度学习")
                await page.screenshot(path=str(screenshots / "project-graph.png"), full_page=True, animations="disabled")
                await page.set_viewport_size({"width": 430, "height": 900})
                await expect(overview.get_by_role("tablist", name="四维浏览")).to_be_visible()
                assert await page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
                await page.screenshot(path=str(screenshots / "projects-mobile.png"), full_page=True, animations="disabled")
                await page.set_viewport_size({"width": 1440, "height": 1000})
                await page.get_by_role("button", name="情报分析", exact=True).click()
                await page.get_by_label("上传外部材料", exact=True).set_input_files(
                    {
                        "name": "外部技术说明.md",
                        "mimeType": "text/markdown",
                        "buffer": "# 外部技术说明\n医学影像方案需要临床数据验证，本资料为功能验收自述。".encode(),
                    }
                )
                await expect(page.get_by_role("button", name="外部技术说明.md", exact=True)).to_be_visible()
                await page.get_by_text("预览和调整报告章节", exact=True).click()
                await page.get_by_label("情报报告模板", exact=True).fill(
                    "# 验收情报报告\n\n## 技术背景\n\n## 来源局限"
                )
                await page.get_by_role("button", name="保存为我的模板", exact=True).click()
                await expect(page.get_by_role("button", name="生成分析报告", exact=True)).to_be_enabled()
                await page.get_by_role("button", name="生成分析报告", exact=True).click()
                pane = page.locator("aside[aria-hidden='false']")
                await expect(pane.get_by_role("heading", name="验收情报报告", exact=True)).to_be_visible()
                await pane.get_by_role("button", name="[1]", exact=True).first.click()
                await expect(pane.get_by_text("外部材料自述 · 原输入版本保留", exact=True)).to_be_visible()
                await page.get_by_role("button", name="关闭检查器", exact=True).click()
                await page.get_by_role("button", name="存量分析", exact=True).click()
                # 成果与回收站统一在检查器中，按报告类别分组。
                await page.get_by_role("button", name="成果与回收站", exact=True).click()
                results = page.locator("aside[aria-hidden='false']")
                await results.get_by_role("tab", name=re.compile("^情报分析报告")).click()
                await expect(results.get_by_text("综合情报报告", exact=True).first).to_be_visible()
                await expect(results.get_by_role("tab", name=re.compile("^资料审查报告"))).to_be_visible()
                await page.get_by_role("button", name="关闭检查器", exact=True).click()

                proposal = WordDocument()
                proposal.add_heading("项目摘要", 1)
                proposal.add_paragraph("摘" * 300)
                proposal.add_heading("项目背景", 1)
                proposal.add_paragraph("采用深度学习开展临床诊疗研究，拟以外部样本验证病灶识别。" * 7 + "补充背景说明内容十四字以上。")
                proposal.add_heading("项目方案", 1)
                proposal.add_paragraph("本方案以临床数据验证模型，条件不足时不能宣称创新。" * 20)
                proposal.add_heading("项目进度计划", 1)
                proposal.add_paragraph("2026年10月至2027年9月完成原型；2028年6月完成验收。")
                data = BytesIO()
                proposal.save(data)
                await page.get_by_role("button", name="资料审查", exact=True).click()
                await expect(page.get_by_role("heading", name="形式审查", exact=True)).to_be_visible()
                await expect(page.get_by_label("形式审查模板")).to_contain_text("通用形式审查示例")
                await expect(page.get_by_text("自然科学基金申请", exact=True)).to_have_count(0)
                await page.get_by_label("上传申请书文件").set_input_files(
                    {
                        "name": "验收申请书.docx",
                        "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        "buffer": data.getvalue(),
                    }
                )
                await expect(page.get_by_label("项目名称", exact=True)).to_be_visible()
                # Validation errors name the field instead of a bare status code.
                await page.get_by_label("项目名称", exact=True).fill("")
                await page.get_by_role("button", name="保存并确认", exact=True).first.click()
                await expect(page.get_by_role("alert").filter(has_text="基本信息保存失败：项目名称")).to_be_visible()
                await page.get_by_label("项目名称", exact=True).fill("影像诊疗验收申请")
                await page.get_by_text("更多字段（基金、类别、年度、学历等，按模板需要填写）", exact=True).first.click()
                await page.get_by_label("正式申报时间", exact=True).first.fill("2026-09-30T10:30")
                await page.get_by_label("项目总预算（万元）", exact=True).fill("280")
                await page.get_by_label("申报单位名单", exact=True).fill("甲大学\n乙医院\n甲大学")
                await page.get_by_role("button", name="保存并确认", exact=True).first.click()
                await expect(page.get_by_text("基本信息已保存。", exact=True)).to_be_visible()
                await expect(page.get_by_label("正式申报时间", exact=True).first).to_have_value("2026-09-30T10:30")
                await expect(page.get_by_text("自动识别 · 待确认").first).to_be_visible()
                await page.get_by_role("button", name="确认并保存章节对应", exact=True).click()
                await expect(page.get_by_text("章节对应 4/4 已确认", exact=True)).to_be_visible()
                await page.screenshot(path=str(screenshots / "formal-workbench.png"), full_page=True, animations="disabled")
                await page.get_by_role("button", name="开始审查", exact=True).click()
                await expect(page.get_by_role("heading", name="形式审查报告", exact=True)).to_be_visible()
                await page.get_by_role("tab", name="正文与大纲", exact=True).click()
                table = page.locator("table").first
                await expect(table.get_by_role("row").filter(has_text="项目背景")).to_contain_text("超出 10")
                await expect(table.get_by_role("row").filter(has_text="项目摘要")).to_contain_text("通过")
                await page.get_by_role("tab", name="信息表与单位", exact=True).click()
                await page.get_by_role("button", name="查看全部结果", exact=True).click()
                await expect(page.get_by_role("article").filter(has_text="申报单位")).to_contain_text("实际值：2")
                await page.screenshot(path=str(screenshots / "formal-review.png"), full_page=True, animations="disabled")

                # Template preview, copy, edit with unsaved protection, and use.
                await page.go_back()
                await expect(page.get_by_label("项目名称", exact=True)).to_have_value("影像诊疗验收申请")
                await page.get_by_role("button", name="预览模板", exact=True).click()
                await expect(page.get_by_role("heading", name="模板预览", exact=True)).to_be_visible()
                await expect(page.get_by_text("项目摘要不超过 300 个字符（含 300）", exact=True).first).to_be_visible()
                await page.get_by_role("button", name="复制编辑", exact=True).click()
                await expect(page.get_by_role("heading", name="编辑模板", exact=True)).to_be_visible()
                await page.get_by_label("上限：项目摘要", exact=True).fill("500")
                await page.get_by_label("要求：项目摘要", exact=True).fill("项目摘要不超过 500 个字符（含 500）")
                await page.get_by_role("link", name="审查模板", exact=True).click()
                await expect(page.get_by_role("alertdialog", name="未保存的修改")).to_be_visible()
                await page.get_by_role("button", name="继续编辑", exact=True).click()
                await page.get_by_label("保存后启用（可用于审查）").check()
                await page.get_by_role("button", name="保存并启用", exact=True).click()
                await expect(page.get_by_text("已保存并启用 v1。", exact=True)).to_be_visible()
                await expect(page.get_by_text("原文要求：项目摘要不超过 300 个字符（含 300）", exact=False)).to_be_visible()
                await page.screenshot(path=str(screenshots / "template-preview.png"), full_page=True, animations="disabled")
                await page.get_by_role("button", name="使用此版本", exact=True).click()
                await expect(page.get_by_label("形式审查模板")).to_contain_text("通用形式审查示例（副本）")
                await expect(page.get_by_label("项目名称", exact=True)).to_have_value("影像诊疗验收申请")

                await page.get_by_role("link", name="专业审查", exact=True).click()
                await expect(page.get_by_role("heading", name="专业审查", exact=True)).to_be_visible()
                await expect(page.get_by_label("专业审查模板")).to_contain_text("通用专业审查")
                # One comparison library is required and preselected; 3 most similar reports by default.
                await expect(page.get_by_label("对照资料库", exact=True)).to_have_value(info.id)
                await expect(page.get_by_label("参考报告篇数", exact=True)).to_have_value("3")
                await page.get_by_role("button", name="开始审查", exact=True).click()
                await expect(page.get_by_role("heading", name="专业审查报告", exact=True)).to_be_visible()
                await page.get_by_role("tab", name="技术专题", exact=True).click()
                await expect(page.get_by_text("创新依据待补充", exact=True)).to_be_visible()
                await expect(page.get_by_text("未完成评价", exact=False).first).to_be_visible()
                await page.get_by_role("button", name="定位这条意见的原文", exact=True).first.click()
                await expect(
                    page.get_by_text("Word 按段落或表格定位，PDF 按原文件页码定位。", exact=True)
                ).to_be_visible()
                await page.get_by_role("button", name="关闭检查器", exact=True).click()
                await page.screenshot(
                    path=str(screenshots / "professional-review.png"), full_page=True, animations="disabled"
                )
                await page.set_viewport_size({"width": 430, "height": 932})
                await page.screenshot(
                    path=str(screenshots / "professional-mobile.png"), full_page=True, animations="disabled"
                )
                assert await page.evaluate("document.documentElement.scrollWidth <= window.innerWidth + 2")
                assert not errors, errors
                await browser.close()
                print(
                    "PASS: five entries, project grouping/chart/relations, intelligence report, generic template preview/copy/edit/use, section mapping, layered formal/professional reports, source reader and narrow layout"
                )
        finally:
            server.should_exit = True
            thread.join(timeout=10)
            listener.close()


if __name__ == "__main__":
    asyncio.run(main())
