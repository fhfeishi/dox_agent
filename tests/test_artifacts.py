"""W3-B: minimal Artifact store, reports compatibility, run linkage and DOCX export."""

import io
import sqlite3

from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from src.agent.config import Settings
from src.agent.corpora import corpus_id_for
from src.artifacts import ArtifactStore
from src.docx_export import markdown_to_docx
from src.knowledge import Document, Knowledge, Page
from src.main import create_app


class FakeGraph:
    def __init__(self, corpus_id):
        self.corpus_id = corpus_id

    async def astream(self, state, **kwargs):
        yield {"event": "sources", "data": [{"doc_id": "d1", "corpus_id": self.corpus_id,
                                             "version": "v1", "title": "报告", "page": 2}]}
        yield {"event": "token", "data": {"text": "回答"}}
        yield {"event": "done", "data": {"ok": True}}


def ready_app(tmp_path):
    root = tmp_path / ".knowledge"
    corpus = root / "fixture"
    corpus.mkdir(parents=True, exist_ok=True)
    settings = Settings(_env_file=None, corpora_root=root, state_dir=tmp_path)
    store = Knowledge(corpus / "datadb" / "knowledge.sqlite3", settings=settings)
    store.put(Document(title="seed", origin="seed", kind="text", parser="text",
                       pages=[Page(number=1, text="正文")]))
    cid = corpus_id_for("fixture")
    return create_app(settings, store, lambda *_: FakeGraph(cid)), cid


def chat_body(cid, run_id):
    return {"run_id": run_id, "session_key": "sess-1", "messages": [{"role": "user", "content": "问题"}],
            "corpus_ids": [cid], "task_id": "task1"}


def run_a_chat(client, cid, run_id):
    assert client.post("/api/chat", json=chat_body(cid, run_id)).status_code == 200


def test_user_can_save_an_answer_as_an_artifact_linked_to_its_run(tmp_path):
    app, cid = ready_app(tmp_path)
    with TestClient(app) as client:
        run_a_chat(client, cid, "chat-run-art-1")
        created = client.post("/api/artifacts", json={
            "type": "answer_snapshot", "run_id": "chat-run-art-1", "title": "癫痫结论",
            "markdown": "回答", "session_key": "sess-1"})
        assert created.status_code == 201
        artifact_id = created.json()["artifact_id"]
        fetched = client.get(f"/api/artifacts/{artifact_id}").json()
        listing = client.get("/api/artifacts?session_key=sess-1").json()
    assert fetched["type"] == "answer_snapshot" and fetched["status"] == "completed"
    assert fetched["run_id"] == "chat-run-art-1" and fetched["current_version"] == 1
    assert fetched["markdown"] == "回答" and fetched["source_verification"] == "verified"
    assert fetched["session_key"] == "sess-1" and fetched["corpus_ids"] == [cid]
    assert fetched["run_available"] is True and listing[0]["run_available"] is True
    # Citations come from the linked run snapshot, not a separate per-artifact store.
    assert fetched["citations"][0]["doc_id"] == "d1" and fetched["citations"][0]["corpus_id"] == cid
    assert artifact_id in {item["artifact_id"] for item in listing}


def test_user_cannot_forge_the_source_of_an_answer_artifact(tmp_path):
    app, cid = ready_app(tmp_path)
    with TestClient(app) as client:
        run_a_chat(client, cid, "chat-run-forgery")
        base = {"type": "answer_snapshot", "run_id": "chat-run-forgery", "markdown": "回答"}
        for changes in (
            {"type": "report"},
            {"session_key": "another-session"},
            {"corpus_ids": ["another-corpus"]},
            {"markdown": "用户改写后冒充原始输出"},
        ):
            response = client.post("/api/artifacts", json={**base, **changes})
            assert response.status_code == 422, (changes, response.json())
        assert client.get("/api/artifacts").json() == []


def test_user_cannot_save_an_incomplete_or_unverified_run_as_an_original_answer(tmp_path):
    app, cid = ready_app(tmp_path)
    with TestClient(app) as client:
        for run_id, status in (("chat-run-running", "running"), ("chat-run-old", "completed")):
            app.state.runs.create(run_id, "fingerprint", run_type="chat", status=status,
                                  session_key="sess-1", effective_corpus_ids=[cid])
            response = client.post("/api/artifacts", json={
                "type": "answer_snapshot", "run_id": run_id, "markdown": "回答"})
            assert response.status_code == 422


def test_user_artifact_uses_null_for_unrecorded_versions_and_retries_preserve_provenance(tmp_path, monkeypatch):
    attempts = 0

    async def generate_once_failed(knowledge, settings, params, *, llm=None, template_content=None,
                                   task_definition=None):
        nonlocal attempts
        attempts += 1
        if attempts == 1:
            raise ValueError("临时失败")
        return "# 重试报告"

    monkeypatch.setattr("src.main.generate_markdown", generate_once_failed)
    app, cid = ready_app(tmp_path)
    with TestClient(app) as client:
        old = app.state.artifacts.create("old-versioned", type="answer_snapshot",
                                         title="旧成果", markdown="旧正文")
        assert old["task_version"] is None and old["template_version"] is None
        template = client.post("/api/templates/custom", json={"source_template_id": "comprehensive"}).json()
        client.put(f"/api/templates/custom/{template['id']}/draft", json={
            "revision": template["revision"], "content": "# 报告\n\n## 章节\n"})
        client.post(f"/api/templates/custom/{template['id']}/publish", json={"revision": 2})
        task = client.post("/api/tasks/custom", json={"source_task_id": "task4"}).json()
        client.put(f"/api/tasks/custom/{task['id']}/draft", json={
            "revision": task["revision"], "goal": "生成报告",
            "report_template_id": template["id"], "report_template_version": 1})
        client.post(f"/api/tasks/custom/{task['id']}/publish", json={"revision": 2})
        request = {"domain": "医疗", "year_from": 2024, "year_to": 2025,
                   "template_id": template["id"], "template_version": 1, "corpus_id": cid,
                   "session_key": "s2", "run_id": "report-provenance-1",
                   "task_id": task["id"], "task_version": 1}
        failed = client.post("/api/reports", json=request)
        assert failed.status_code == 422
        retried = client.post("/api/reports", json=request)
        assert retried.status_code == 201
        run = client.get("/api/runs/report-provenance-1").json()
        assert run["params"]["template_id"] == template["id"]
        assert run["params"]["template_version"] == 1
        assert run["task_id"] == task["id"] and run["task_version"] == 1
        artifacts = client.get("/api/artifacts?session_key=s2").json()
        assert len(artifacts) == 1
        artifact = artifacts[0]
        assert artifact["current_version"] == 2 and artifact["status"] == "completed"
        assert artifact["template_id"] == template["id"]
        assert artifact["template_version"] == 1
        assert artifact["task_id"] == task["id"] and artifact["task_version"] == 1
        app.state.artifacts.create("old-artifact", type="answer_snapshot", title="旧成果",
                                   markdown="旧正文", run_id="chat-run-old")
        old = client.get("/api/artifacts/old-artifact").json()
        assert old["source_verification"] == "unverified"


def test_new_artifact_version_does_not_overwrite_the_old_version(tmp_path):
    app, _ = ready_app(tmp_path)
    with TestClient(app):
        store = app.state.artifacts
        store.create("art-1", type="answer_snapshot", title="v1 标题", markdown="# 第一版")
        store.add_version("art-1", markdown="# 第二版")
        current = store.get("art-1")
        first = store.get_version("art-1", 1)
    assert current["current_version"] == 2 and current["markdown"] == "# 第二版"
    assert first["markdown"] == "# 第一版"


def test_user_existing_artifact_database_migrates_without_claiming_old_output_verified(tmp_path):
    path = tmp_path / "artifacts.sqlite3"
    with sqlite3.connect(path) as db:
        db.execute("CREATE TABLE artifact_versions (artifact_id TEXT NOT NULL, version INTEGER NOT NULL, "
                   "created_at TEXT NOT NULL, markdown TEXT NOT NULL DEFAULT '', "
                   "citations TEXT NOT NULL DEFAULT '[]', PRIMARY KEY (artifact_id, version))")
    store = ArtifactStore(path)
    old = store.create("old", type="answer_snapshot", title="旧成果", markdown="旧输出")
    assert old["source_verification"] == "unverified"
    revised = store.add_version("old", markdown="修改后")
    assert revised["version"] == 2 and revised["source_verification"] == "user_modified"


def test_user_edit_creates_a_persistent_revision_with_versioned_export(tmp_path):
    app, cid = ready_app(tmp_path)
    with TestClient(app) as client:
        run_a_chat(client, cid, "chat-run-version")
        created = client.post("/api/artifacts", json={
            "run_id": "chat-run-version", "markdown": "回答"}).json()
        artifact_id = created["artifact_id"]
        edited = client.post(f"/api/artifacts/{artifact_id}/versions", json={
            "markdown": "# 用户修订\n\n补充说明", "status": "draft"})
        assert edited.status_code == 201
        assert edited.json()["version"] == 2 and edited.json()["status"] == "draft"
        assert edited.json()["source_verification"] == "user_modified"
        old = client.get(f"/api/artifacts/{artifact_id}?version=1").json()
        assert old["markdown"] == "回答" and old["source_verification"] == "verified"
        versions = client.get(f"/api/artifacts/{artifact_id}/versions").json()
        assert [item["version"] for item in versions] == [1, 2]
        old_export = client.get(f"/api/artifacts/{artifact_id}/export?format=md&version=1")
        new_export = client.get(f"/api/artifacts/{artifact_id}/export?format=md&version=2")
        assert old_export.text == "回答" and new_export.text.startswith("# 用户修订")
        assert client.post(f"/api/artifacts/{artifact_id}/versions", json={
            "markdown": "invalid", "status": "generating"}).status_code == 422
        assert client.get(f"/api/artifacts/{artifact_id}?version=99").status_code == 404

    restarted, _ = ready_app(tmp_path)
    with TestClient(restarted) as client:
        latest = client.get(f"/api/artifacts/{artifact_id}").json()
        first = client.get(f"/api/artifacts/{artifact_id}?version=1").json()
    assert latest["version"] == 2 and latest["status"] == "draft"
    assert first["version"] == 1 and first["markdown"] == "回答"


def test_artifact_list_distinguishes_omitted_and_empty_session_filters(tmp_path):
    app, _ = ready_app(tmp_path)
    with TestClient(app) as client:
        app.state.artifacts.create("a-empty", type="answer_snapshot", title="空会话", markdown="# a")
        app.state.artifacts.create("a-s1", type="answer_snapshot", title="会话一", markdown="# b",
                                   session_key="s1")
        global_items = client.get("/api/artifacts").json()
        empty_items = client.get("/api/artifacts?session_key=").json()
        s1_items = client.get("/api/artifacts?session_key=s1").json()
    assert {item["artifact_id"] for item in global_items} >= {"a-empty", "a-s1"}
    assert {item["artifact_id"] for item in empty_items} == {"a-empty"}
    assert {item["artifact_id"] for item in s1_items} == {"a-s1"}


def test_legacy_report_is_readable_as_an_artifact_without_rewriting_it(tmp_path):
    app, _ = ready_app(tmp_path)
    with TestClient(app) as client:
        app.state.reports.save("legacy", {"template_id": "achievements", "domain": "旧领域"},
                               "# 旧报告", session_key="s1", run_id="", corpus_id="")
        listing = client.get("/api/artifacts?session_key=s1").json()
        detail = client.get("/api/artifacts/report:legacy").json()
        still_there = app.state.reports.get("legacy")
    assert any(item["artifact_id"] == "report:legacy" and item["legacy"] for item in listing)
    assert detail["markdown"] == "# 旧报告" and detail["type"] == "report"
    assert detail["run_id"] == "" and detail["corpus_ids"] == []  # "未记录", never inferred
    assert detail["run_available"] is False
    assert still_there["markdown"] == "# 旧报告"


def test_report_generation_creates_a_linked_artifact(tmp_path, monkeypatch):
    async def fake_generate(knowledge, settings, params, *, llm=None):
        return "# 报告"

    monkeypatch.setattr("src.main.generate_markdown", fake_generate)
    app, cid = ready_app(tmp_path)
    with TestClient(app) as client:
        created = client.post("/api/reports", json={
            "domain": "医疗", "year_from": 2024, "year_to": 2025, "template_id": "comprehensive",
            "corpus_id": cid, "session_key": "sess-1", "run_id": "report-run-art-1",
            "parent_run_id": "intake-run-1"})
        assert created.status_code == 201
        listing = client.get("/api/artifacts?session_key=sess-1").json()
    reports = [item for item in listing if item["type"] == "report" and item["run_id"] == "report-run-art-1"]
    assert len(reports) == 1 and reports[0]["current_version"] == 1
    # The compatible legacy entry is suppressed because the run is already linked by an artifact.
    assert all(item["artifact_id"] != "report:" + created.json()["report_id"] for item in listing)


def test_user_report_retry_repairs_a_failed_artifact_link(tmp_path, monkeypatch):
    async def fake_generate(knowledge, settings, params, *, llm=None):
        return "# 报告"

    monkeypatch.setattr("src.main.generate_markdown", fake_generate)
    app, cid = ready_app(tmp_path)
    with TestClient(app) as client:
        original = app.state.artifacts.ensure_report
        attempts = 0

        def fail_once(report):
            nonlocal attempts
            attempts += 1
            if attempts == 1:
                raise OSError("temporary artifact write error")
            return original(report)

        monkeypatch.setattr(app.state.artifacts, "ensure_report", fail_once)
        request = {"domain": "医疗", "year_from": 2024, "year_to": 2025,
                   "template_id": "comprehensive", "corpus_id": cid,
                   "session_key": "sess-1", "run_id": "report-run-repair"}
        first = client.post("/api/reports", json=request)
        assert first.status_code == 201
        second = client.post("/api/reports", json=request)
        assert second.status_code == 200 and second.json()["idempotent"] is True
        listing = client.get("/api/artifacts?session_key=sess-1").json()
    assert attempts == 2
    assert len([item for item in listing if item["type"] == "report"]) == 1
    assert listing[0]["run_id"] == "report-run-repair" and not listing[0].get("legacy")


def test_user_failed_report_is_visible_and_retry_keeps_the_failed_version(tmp_path, monkeypatch):
    attempts = 0

    async def generate_once_failed(knowledge, settings, params, *, llm=None):
        nonlocal attempts
        attempts += 1
        if attempts == 1:
            raise ValueError("资料不足")
        return "# 重试成功\n\n来源 [1]"

    monkeypatch.setattr("src.main.generate_markdown", generate_once_failed)
    app, cid = ready_app(tmp_path)
    request = {"domain": "医疗", "year_from": 2024, "year_to": 2025,
               "template_id": "comprehensive", "corpus_id": cid,
               "session_key": "sess-1", "run_id": "report-run-failure"}
    with TestClient(app) as client:
        failed = client.post("/api/reports", json=request)
        assert failed.status_code == 422
        failed_items = client.get("/api/artifacts?status=failed").json()
        assert len(failed_items) == 1 and failed_items[0]["run_id"] == "report-run-failure"
        artifact_id = failed_items[0]["artifact_id"]
        assert failed_items[0]["fail_reason"] == "资料不足"
        retry = client.post("/api/reports", json=request)
        assert retry.status_code == 201
        current = client.get(f"/api/artifacts/{artifact_id}").json()
        first = client.get(f"/api/artifacts/{artifact_id}?version=1").json()
        assert current["status"] == "completed" and current["version"] == 2
        assert current["source_verification"] == "verified" and current["markdown"].startswith("# 重试成功")
        assert first["status"] == "failed" and first["markdown"] == ""
        assert first["fail_reason"] == "资料不足"
    assert len([item for item in client.get("/api/artifacts").json()
                    if item["run_id"] == "report-run-failure"]) == 1


def test_user_sees_a_safe_failed_report_status_without_provider_details(tmp_path, monkeypatch):
    async def generate_failed(knowledge, settings, params, *, llm=None):
        raise RuntimeError("private provider token")

    monkeypatch.setattr("src.main.generate_markdown", generate_failed)
    app, cid = ready_app(tmp_path)
    with TestClient(app) as client:
        response = client.post("/api/reports", json={
            "domain": "医疗", "year_from": 2024, "year_to": 2025,
            "template_id": "comprehensive", "corpus_id": cid,
            "session_key": "sess-1", "run_id": "report-run-provider-failure"})
        assert response.status_code == 500
        assert "private provider token" not in response.text
        failed = client.get("/api/artifacts?status=failed").json()
        assert len(failed) == 1 and failed[0]["fail_reason"] == "报告生成失败"
        assert client.get("/api/runs/report-run-provider-failure").json()["status"] == "failed"


def test_user_global_report_list_keeps_older_legacy_reports_beyond_one_page(tmp_path):
    app, _ = ready_app(tmp_path)
    with TestClient(app) as client:
        app.state.reports.save("legacy-old", {"domain": "历史报告"}, "# 旧", session_key="")
        for index in range(102):
            run_id = f"linked-run-{index}"
            app.state.reports.save(f"report-{index}", {"domain": f"R{index}"}, "# 正文",
                                   session_key="s1", run_id=run_id)
            app.state.artifacts.create(f"artifact-{index}", type="report", title=f"R{index}",
                                       markdown="# 正文", session_key="s1", run_id=run_id)
        listing = client.get("/api/artifacts?limit=150&type=report").json()
    assert len(listing) == 103
    assert len({item["run_id"] for item in listing if item["run_id"]}) == 102
    assert len([item for item in listing if item["artifact_id"] == "report:legacy-old"]) == 1


def test_docx_export_is_real_ooxml_with_headings_and_tables(tmp_path):
    markdown = ("# 标题\n\n段落文字 [1]\n\n| 名称 | 值 |\n| --- | --- |\n| A | 1 |\n\n"
                "来源：[1] 报告 A（文档 d1）\n")
    data = markdown_to_docx(markdown)
    assert data[:2] == b"PK"  # real OOXML zip, not HTML
    document = DocxDocument(io.BytesIO(data))
    assert any(paragraph.text == "标题" for paragraph in document.paragraphs)
    assert len(document.tables) == 1 and document.tables[0].rows[0].cells[0].text == "名称"
    assert any("[1] 报告 A（文档 d1）" in paragraph.text for paragraph in document.paragraphs)

    app, _ = ready_app(tmp_path)
    with TestClient(app) as client:
        app.state.artifacts.create("art-docx", type="report", title="导出", markdown=markdown)
        exported = client.get("/api/artifacts/art-docx/export?format=docx")
        markdown_export = client.get("/api/artifacts/art-docx/export?format=md")
    assert exported.status_code == 200 and exported.content[:2] == b"PK"
    assert "wordprocessingml" in exported.headers["content-type"]
    assert markdown_export.text == markdown


def test_saving_an_artifact_for_an_unknown_run_is_rejected(tmp_path):
    app, _ = ready_app(tmp_path)
    with TestClient(app) as client:
        response = client.post("/api/artifacts", json={
            "type": "answer_snapshot", "run_id": "never-persisted", "markdown": "# x"})
    assert response.status_code == 422
