"""Replaceable ingestion functions: local mineru (PDF), docx, Crawl4AI or Firecrawl."""

import asyncio
import hashlib
import json
import os
import shutil
import subprocess
from pathlib import Path
from urllib.parse import urlsplit

import httpx

from .agent.config import Settings
from .knowledge import Document, Page


def _first_file(directory: Path, names: list[str]) -> Path | None:
    for name in names:
        candidate = directory / name
        if candidate.is_file():
            return candidate
    for pattern in names:
        matches = sorted(directory.glob(pattern))
        if matches:
            return matches[0]
    return None


def _block_text(block: dict) -> str:
    if isinstance(block.get("text"), str):
        return block["text"]
    content = block.get("content")
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        return "".join(str(item.get("content", "")) if isinstance(item, dict) else str(item) for item in content)
    return ""


def read_mineru_output(out_dir: Path) -> tuple[str, list[Page]]:
    """K13: read a mineru output dir into ``(markdown, pages)``.

    v4 emits one ``*.md`` (base64 images) plus a middle JSON with ``pages[].page_idx``;
    classic emits ``full.md`` plus ``*_content_list.json`` blocks with ``page_idx``.
    Page numbers come from the JSON; without it the markdown is a single page.
    """
    markdown_path = _first_file(out_dir, ["full.md", "*.md"])
    markdown = markdown_path.read_text(encoding="utf-8") if markdown_path else ""

    by_page: dict[int, list[str]] = {}
    for candidate in sorted(out_dir.glob("*.json")):
        try:
            data = json.loads(candidate.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        grouped: list[tuple[int, list]] = []
        if isinstance(data, dict) and isinstance(data.get("pages"), list):
            grouped = [(page.get("page_idx", index), page.get("blocks", [])) for index, page in enumerate(data["pages"])]
        elif isinstance(data, list):
            pages: dict[int, list] = {}
            for block in data:
                if isinstance(block, dict):
                    pages.setdefault(int(block.get("page_idx", 0)), []).append(block)
            grouped = list(pages.items())
        for page_idx, page_blocks in grouped:
            texts = [text for text in (_block_text(block) for block in page_blocks if isinstance(block, dict)) if text.strip()]
            if texts:
                by_page.setdefault(int(page_idx), []).extend(texts)
        if by_page:
            break

    pages = [Page(number=index + 1, text="\n".join(texts)) for index, texts in sorted(by_page.items())]
    if not pages and markdown.strip():
        pages = [Page(number=1, text=markdown)]
    return markdown, pages


def parse_pdf_pages(path: Path, settings: Settings, out_dir: Path) -> list[Page]:
    """K13: run mineru into ``out_dir`` and return per-page text (page_idx + 1)."""
    shutil.rmtree(out_dir, ignore_errors=True)
    out_dir.mkdir(parents=True, exist_ok=True)
    command = settings.mineru_cmd.format(pdf=str(path), out=str(out_dir))
    environment = {**os.environ}
    if settings.mineru_home:
        environment["MINERU_HOME"] = str(settings.mineru_home)
    try:
        subprocess.run(command, shell=True, check=True, env=environment,
                       timeout=settings.mineru_timeout, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
    except subprocess.TimeoutExpired as exc:
        raise ValueError("mineru 解析超时（MINERU_TIMEOUT）") from exc
    except subprocess.CalledProcessError as exc:
        tail = (exc.output or b"")[-500:].decode("utf-8", "ignore")
        raise ValueError("mineru 解析失败：" + tail) from exc
    _, pages = read_mineru_output(out_dir)
    if not pages:
        raise ValueError("mineru 输出为空，未入库")
    return pages


def parse_file(path: Path, settings: Settings, *, parsed_dir: Path | None = None) -> Document:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        pages = [Page(number=1, text=path.read_text(encoding="utf-8-sig"))]
        parser = "utf8"
    elif suffix == ".pdf":
        out_dir = parsed_dir or path.parent / ".mineru" / path.stem
        pages = parse_pdf_pages(path, settings, out_dir)
        parser = "mineru"
    elif suffix == ".docx":
        # K8: offline Word parsing (paragraphs + tables as text).
        import docx

        document = docx.Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        pages = [Page(number=1, text="\n".join(parts))]
        parser = "python-docx/" + getattr(docx, "__version__", "1")
    else:
        raise ValueError("仅支持 txt、md、pdf、docx")
    return Document(
        title=path.stem,
        origin=str(path.resolve()),
        kind="pdf" if suffix == ".pdf" else "text",
        parser=parser,
        pages=pages,
    )


def session_for(url: str, settings: Settings) -> dict:
    if not settings.web_sessions_file:
        return {}
    sessions = json.loads(settings.web_sessions_file.read_text(encoding="utf-8"))
    return sessions.get(urlsplit(url).hostname, {})


async def crawl4ai_page(url: str, session: dict) -> tuple[str, str]:
    from crawl4ai import AsyncWebCrawler, BrowserConfig, CacheMode, CrawlerRunConfig

    browser = BrowserConfig(headless=True, verbose=False, storage_state=session.get("storage_state"))
    config = CrawlerRunConfig(cache_mode=CacheMode.BYPASS, verbose=False, page_timeout=60000)
    async with AsyncWebCrawler(config=browser) as crawler:
        result = await crawler.arun(url=url, config=config)
    if not result.success or (result.status_code or 200) >= 400:
        raise ValueError("网页抓取失败，可能需要登录、重新授权或人工处理")
    return (result.metadata or {}).get("title") or url, result.markdown.raw_markdown


async def firecrawl_page(url: str, session: dict, settings: Settings) -> tuple[str, str]:
    if not settings.firecrawl_api_key:
        raise ValueError("请配置 FIRECRAWL_API_KEY")
    async with httpx.AsyncClient(timeout=90) as client:
        response = await client.post(
            settings.firecrawl_base_url.rstrip("/") + "/v2/scrape",
            headers={"Authorization": "Bearer " + settings.firecrawl_api_key.get_secret_value()},
            json={
                "url": url,
                "formats": ["markdown"],
                "onlyMainContent": True,
                "headers": session.get("headers", {}),
            },
        )
    if response.status_code >= 400:
        raise ValueError(f"Firecrawl 请求失败（HTTP {response.status_code}）")
    data = response.json()
    if not data.get("success"):
        raise ValueError("Firecrawl 未能提取页面")
    result = data["data"]
    return result.get("metadata", {}).get("title", url), result.get("markdown", "")


async def parse_web(url: str, settings: Settings) -> Document:
    parsed = urlsplit(url)
    if parsed.scheme not in {"https", "http"} or not parsed.hostname or parsed.username:
        raise ValueError("请输入不含账号密码的 HTTP(S) URL")
    session = session_for(url, settings)
    if settings.web_provider == "crawl4ai":
        title, text = await asyncio.wait_for(crawl4ai_page(url, session), timeout=100)
    elif settings.web_provider == "firecrawl":
        title, text = await firecrawl_page(url, session, settings)
    else:
        raise ValueError("WEB_PROVIDER 必须为 crawl4ai 或 firecrawl")
    if not text.strip():
        raise ValueError("网页正文为空，请检查登录状态或页面加载情况")
    if len(text) > 1_000_000:
        raise ValueError("页面超过 100 万字符，请缩小抓取范围")
    return Document(
        title=title, origin=url, kind="web", parser=settings.web_provider, pages=[Page(number=1, text=text)]
    )


def import_defaults(knowledge, settings: Settings, *, root: Path | None = None, exclude: list[Path] = (), force: bool = False, parsed_root: Path | None = None) -> dict:
    """K1: incremental local import backed by the `files` manifest.

    With `root` (a corpus ``source/`` dir) only that corpus is collected; the legacy path
    keeps the old dual-root collection. Unchanged files are skipped; changed/new files are
    re-parsed; source files that disappeared are removed from the manifest and from `docs`
    (so BM25 stops returning them; dense drops stale ids on its next sync). An empty manifest
    is the first backfill and re-parses everything once. ``parsed_root`` caches mineru output
    per file (K13).
    """
    excluded = [item.resolve() for item in exclude]
    manifest = knowledge.files()
    current: dict[str, Path] = {}
    for path in collect_sources(root, settings):
        resolved = path.resolve()
        if any(resolved == item or item in resolved.parents for item in excluded):
            continue
        base = Path(root) if root is not None else source_base(path, settings)
        try:
            rel = resolved.relative_to(base.resolve()).as_posix()
        except ValueError:
            rel = path.name
        current[rel] = path

    added = updated = skipped = 0
    imported, errors = [], []
    for rel, path in current.items():
        prev = manifest.get(rel)
        try:
            stat = path.stat()
            if not force and prev and prev["status"] == "indexed" and prev["size"] == stat.st_size and prev["mtime_ns"] == stat.st_mtime_ns:
                skipped += 1
                continue
            digest = sha256_file(path)
            if not force and prev and prev["status"] == "indexed" and prev["sha256"] == digest:
                knowledge.record_file(rel, stat.st_size, stat.st_mtime_ns, digest, prev["doc_id"], "indexed")
                skipped += 1
                continue
            result = knowledge.put(parse_file(path, settings, parsed_dir=(parsed_root / rel) if parsed_root else None))
            knowledge.record_file(rel, stat.st_size, stat.st_mtime_ns, digest, result["doc_id"], "indexed")
            imported.append(result)
            updated += 1 if prev else 0
            added += 0 if prev else 1
        except Exception as exc:  # noqa: BLE001 - retain other sources on parser failure
            errors.append({"source": path.name, "error": type(exc).__name__})
            try:
                stat = path.stat()
                knowledge.record_file(rel, stat.st_size, stat.st_mtime_ns, (prev or {}).get("sha256", ""),
                                      (prev or {}).get("doc_id"), "error")
            except OSError:
                pass

    deleted = 0
    for rel in list(manifest):
        if rel in current or manifest[rel]["status"] == "removed":
            continue
        knowledge.drop_file(rel)
        deleted += 1
    return {"imported": imported, "errors": errors, "added": added, "updated": updated,
            "skipped": skipped, "deleted": deleted, "scanned": len(current)}


def source_base(path: Path, settings: Settings) -> Path:
    """Legacy collection spans knowledge_root/text_root; pick the root that contains the file."""
    for candidate in (settings.knowledge_root, settings.text_root):
        try:
            path.resolve().relative_to(candidate.resolve())
            return candidate
        except ValueError:
            continue
    return path.parent


def collect_sources(root: Path | None, settings: Settings) -> list[Path]:
    if root is not None:
        base = Path(root)
        return sorted(p for p in base.rglob("*") if p.is_file() and p.suffix.lower() in {'.pdf', '.txt', '.md'})
    return sorted(
        set(settings.text_root.rglob("*.txt"))
        | set(settings.text_root.rglob("*.md"))
        | {p for p in settings.knowledge_root.rglob("*") if p.is_file() and p.suffix.lower() == ".pdf"}
    )


def sha256_file(path: Path, chunk: int = 1 << 20) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(chunk), b""):
            digest.update(block)
    return digest.hexdigest()
