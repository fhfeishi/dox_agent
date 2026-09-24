"""W3-B: real OOXML (.docx) export from report/answer Markdown (frozen spec §16.15.4).

The prior "Word" export was Word-compatible HTML (``.doc``); this is a genuine ``.docx`` built
with python-docx so headings, tables and paragraphs survive a round trip.
"""

import io
import re

from docx import Document

_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_TABLE_SEPARATOR = re.compile(r"^\s*\|?[\s:|-]+\|?\s*$")


def _split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _plain(text: str) -> str:
    return re.sub(r"(?:\*\*|__|[*_`])", "", text).strip()


def markdown_to_docx(markdown: str) -> bytes:
    """Render Markdown headings, tables and paragraphs into a real ``.docx`` byte string."""
    document = Document()
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        heading = _HEADING.match(line)
        if heading:
            document.add_heading(_plain(heading.group(2)), level=min(len(heading.group(1)), 6))
            index += 1
            continue
        if "|" in line and index + 1 < len(lines) and _TABLE_SEPARATOR.match(lines[index + 1]):
            header = _split_row(line)
            index += 2
            rows = []
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                rows.append(_split_row(lines[index]))
                index += 1
            if header:
                table = document.add_table(rows=1, cols=len(header))
                for column, text in enumerate(header):
                    table.rows[0].cells[column].text = _plain(text)
                for row in rows:
                    cells = table.add_row().cells
                    for column, text in enumerate(row[:len(header)]):
                        cells[column].text = _plain(text)
            continue
        if line.strip():
            document.add_paragraph(_plain(line))
        index += 1
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
